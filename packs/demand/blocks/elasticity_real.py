"""elasticity_real@1.0 — kind `model`. Real per-SKU price elasticity for a
brand that has an actual shipped ds-models run, read straight from that
run's elasticity workbook (`packs/demand/data/real/`) — the production
DML-residualized, hierarchical-Bayesian-shrunk, causal-forest-cross-checked,
PSM-DiD-calibrated numbers, not the `elasticity_dml` demo stand-in.

No panel input: unlike `elasticity_dml`, the model already ran inside
ds-models — this block only reads its output — so there is nothing upstream
to consume and the elasticity_surface@v1 it produces still goes through the
same mandatory `post_model_checks` gate as any other model-kind block.
"""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.execute.adapter import BlockContext, build_envelope
from core.execute.envelope import Envelope

_DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "real"
_SHEET = "2. SKU x National"
# rows[0..3] are the workbook's own merged title/section-header block (see
# CODEBASE.md §3.7's "workbook contract" for the shipped layout); rows[4] is
# the real column header row.
_HEADER_ROW_INDEX = 4

# brand slug (core.classify.entities.KNOWN_BRANDS) -> (workbook filename under
# _DATA_DIR, own-brand label as it appears in that workbook's own Brand
# column — every other row is an admitted pooled competitor, per ds-models'
# multi-focal pooling design; §2.6 of CODEBASE.md).
_REGISTRY: dict[str, tuple[str, str]] = {
    "gozero": ("gozero_sticks_price_elasticity_pooled_2026-08-05_072359.xlsx", "go zero"),
    "haleon": ("haleon_sensodyne_oral_care_price_elasticity_pooled_2026-08-05_095359.xlsx", "sensodyne"),
}

# Exported so panel_builder (the synthetic pipeline's first node) can refuse a
# real brand outright: the planner's L2 template cache keys on entity *keys*,
# not values (core.classify.classifier.normalized_hash, deliberately, for
# unrelated reasons), so a demo-shaped cached template could otherwise be
# rebound with brand="gozero"/"haleon" and silently execute the dummy
# synthetic model for a brand that has a real ds-models run — fail loud
# instead of fail silent.
REAL_BRANDS = frozenset(_REGISTRY)


def _load_own_brand_rows(brand: str) -> list[dict[str, Any]]:
    try:
        filename, own_brand_label = _REGISTRY[brand]
    except KeyError:
        raise ValueError(
            f"no real ds-models elasticity workbook registered for brand {brand!r}; "
            f"known: {sorted(_REGISTRY)}"
        ) from None

    workbook = load_workbook(_DATA_DIR / filename, read_only=True, data_only=True)
    try:
        sheet = workbook[_SHEET]
        all_rows = list(sheet.iter_rows(values_only=True))
        header = all_rows[_HEADER_ROW_INDEX]
        col = {name: idx for idx, name in enumerate(header) if name}

        rows: list[dict[str, Any]] = []
        for raw in all_rows[_HEADER_ROW_INDEX + 1 :]:
            sku_id = raw[col["SKU ID"]]
            if sku_id is None:
                continue
            if str(raw[col["Brand"]]).strip().lower() != own_brand_label:
                continue  # an admitted pooled competitor SKU, not this brand's own
            rows.append(
                {
                    "sku": str(sku_id),
                    "sku_name": raw[col["SKU Name"]],
                    "category": raw[col["Category"]],
                    "pack_tier": raw[col["Pack Tier"]],
                    "elasticity": float(raw[col["Elasticity"]]),
                    "ci_lower": float(raw[col["CI Low"]]) if raw[col["CI Low"]] is not None else None,
                    "ci_upper": float(raw[col["CI High"]]) if raw[col["CI High"]] is not None else None,
                    "confidence": raw[col["Confidence"]],
                    "method": "ds_models_pooled_dml_bayes_causal_forest_psm_did_calibrated",
                }
            )
        return rows
    finally:
        workbook.close()


def run(params: dict[str, Any], inputs: dict[str, Envelope], ctx: BlockContext) -> dict[str, Envelope]:
    brand = str(params["brand"]).strip().lower()
    surface_rows = _load_own_brand_rows(brand)

    ref = ctx.storage.put("surface", surface_rows)
    mean_elasticity = (
        round(sum(row["elasticity"] for row in surface_rows) / len(surface_rows), 3) if surface_rows else 0.0
    )
    summary = (
        f"{len(surface_rows)} SKU elasticities for {brand} — real ds-models run "
        f"(pooled DML + Bayesian shrinkage + causal-forest cross-check, PSM-DiD calibrated); "
        f"mean={mean_elasticity}"
    )

    envelope = build_envelope(
        block_version=ctx.block_version,
        params=params,
        inputs=inputs,
        summary=summary,
        cols=["sku", "sku_name", "category", "pack_tier", "elasticity", "ci_lower", "ci_upper", "confidence", "method"],
        sample=surface_rows[:5],
        ref=ref,
        rows=len(surface_rows),
    )
    return {"surface": envelope}
